# -*- coding: utf-8 -*-
"""版式引擎基础件：设计令牌、几何换算、文本与表格组件。

被 deck_layouts.py 使用，不单独运行。画布固定 1280×720，所有尺寸单位为
96dpi 像素（与设计稿一致），落盘时换算成 EMU / pt。

本文件只负责「怎么画」，不负责「画什么」。
"""
import math

from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ---------- 设计令牌 ----------

PALETTE = {
    'navy':   RGBColor(0x06, 0x1F, 0x32),
    'blue':   RGBColor(0x00, 0x19, 0x65),
    'cyan':   RGBColor(0x08, 0xA6, 0xF6),
    'lblue':  RGBColor(0x99, 0xBD, 0xED),
    'red':    RGBColor(0x5A, 0x0E, 0x13),
    'green':  RGBColor(0x17, 0x7B, 0x57),
    'orange': RGBColor(0xF4, 0x7B, 0x39),
    'gray':   RGBColor(0x80, 0x80, 0x80),
    'lgray':  RGBColor(0xE7, 0xE6, 0xE6),
    'xlgray': RGBColor(0xF5, 0xF5, 0xF5),
    'dgray':  RGBColor(0x44, 0x54, 0x6A),
    'white':  RGBColor(0xFF, 0xFF, 0xFF),
    'black':  RGBColor(0x00, 0x00, 0x00),
    'navy2':  RGBColor(0x2A, 0x44, 0x5A),
}

EA = '微软雅黑'
LATIN = 'Arial'
CANVAS_W, CANVAS_H = 1280, 720
MARGIN = 60
CONTENT_W = CANVAS_W - MARGIN * 2


def color(name):
    """把 brief 里的颜色名换成 RGBColor；已是 RGBColor 则原样返回。"""
    if isinstance(name, RGBColor):
        return name
    try:
        return PALETTE[name]
    except KeyError:
        raise KeyError('未知颜色 %r，可用：%s' % (name, ', '.join(sorted(PALETTE))))


def px(v):
    return Emu(int(round(v * 914400 / 96)))


def pt(v):
    return Pt(v * 0.75)


# ---------- 绘制登记表（供交付自检使用） ----------

_REG = []


def reset_registry():
    _REG.clear()


def registry():
    """返回当前页的绘制登记表：[(order, kind, x, y, w, h, meta), ...]"""
    return list(_REG)


def _register(kind, x, y, w, h, meta):
    _REG.append((len(_REG), kind, x, y, w, h, meta))


# ---------- 文本 ----------


def _set_font(run, size_px, col, bold=False):
    f = run.font
    f.size = pt(size_px)
    f.bold = bold
    f.color.rgb = col
    f.name = LATIN
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', EA)


def text_width(text, size):
    """粗略估算单行宽度：全角按 1 个字号，半角按 0.55 个字号。"""
    w = 0.0
    for ch in text:
        w += size if ord(ch) > 0x2E7F else size * 0.55
    return w


def est_height(lines, box_w):
    """估算一组段落撑开后的总高度，用于溢出预警。"""
    total = 0.0
    for ln in lines:
        size = ln['size']
        ls = ln.get('line_spacing', 1.35)
        usable = max(box_w - 8, 40)
        n = max(1, int(math.ceil(text_width(ln['text'], size) / usable)))
        total += n * size * 1.2 * ls
        total += ln.get('space_before', 0) * 0.75
    return total


def textbox(slide, x, y, w, h, lines, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT,
            margins=(0, 0, 0, 0), wrap=True):
    """lines: [{'text','size','color','bold','line_spacing','space_before','align'}, ...]"""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [px(m) for m in margins]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get('align', align)
        p.line_spacing = ln.get('line_spacing', 1.35)
        if 'space_before' in ln:
            p.space_before = px(ln['space_before'])
        if 'space_after' in ln:
            p.space_after = px(ln['space_after'])
        r = p.add_run()
        r.text = ln['text']
        _set_font(r, ln['size'], color(ln['color']), ln.get('bold', False))
    _register('text', x, y, w, h, {'est': est_height(lines, w), 'preview': lines[0]['text'][:18]})
    return tb


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(x), px(y), px(w), px(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = color(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = color(line)
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    _register('rect', x, y, w, h, {'fill': fill})
    return sh


# ---------- 组件 ----------


def header(slide, title, page, total, source=None, two_line=False):
    """母版：Action Title + 分割线 + 页脚。返回正文起始 y。"""
    h = 96 if two_line else 54
    textbox(slide, MARGIN, 44, CONTENT_W, h,
            [{'text': title, 'size': 36, 'color': 'navy', 'bold': True, 'line_spacing': 1.25}])
    div_y = 44 + h + 8
    rect(slide, MARGIN, div_y, CONTENT_W, 1.5, fill='black')
    if source:
        textbox(slide, MARGIN, 662, 900, 34,
                [{'text': source, 'size': 13, 'color': 'gray', 'line_spacing': 1.3}])
    textbox(slide, 1100, 662, 120, 34,
            [{'text': '%02d/%d' % (page, total), 'size': 14, 'color': 'gray',
              'align': PP_ALIGN.RIGHT}])
    return div_y + 14


def table(slide, x, y, col_ws, rows, row_h=52, head_h=50, body_size=20, head_size=19,
          aligns=None, zebra=True):
    """用矩形拼表：主题表样式不可控，拼出来才能锁死行高与配色。

    rows[0] 为表头；单元格可为字符串或 {'t','c','b'} 字典。
    返回表格底边 y。
    """
    cur_y = y
    for ri, row in enumerate(rows):
        is_head = (ri == 0)
        rh = head_h if is_head else row_h
        cur_x = x
        for ci, cell in enumerate(row):
            if isinstance(cell, dict):
                txt = cell.get('t', '')
                cname = cell.get('c')
                force_bold = cell.get('b')
            else:
                txt, cname, force_bold = cell, None, None
            if is_head:
                bg, tc, bold, fs = 'navy', 'white', True, head_size
            else:
                bg = 'xlgray' if (zebra and ri % 2 == 0) else 'white'
                tc = cname or 'dgray'
                bold = cname is not None if force_bold is None else force_bold
                fs = body_size
            rect(slide, cur_x, cur_y, col_ws[ci], rh, fill=bg, line='lgray', line_w=0.75)
            al = aligns[ci] if aligns else PP_ALIGN.LEFT
            pad = 16 if al == PP_ALIGN.LEFT else 10
            textbox(slide, cur_x + pad, cur_y, col_ws[ci] - pad * 2, rh,
                    [{'text': txt, 'size': fs, 'color': tc, 'bold': bold,
                      'align': al, 'line_spacing': 1.15}],
                    anchor=MSO_ANCHOR.MIDDLE)
            cur_x += col_ws[ci]
        cur_y += rh
    return cur_y


def bignum(slide, x, y, w, h, value, unit, value_size=64, unit_size=20,
           value_color='navy', unit_color='gray', align=PP_ALIGN.LEFT):
    """大数字 + 单位，主视觉元素。"""
    textbox(slide, x, y, w, h,
            [{'text': value, 'size': value_size, 'color': value_color, 'bold': True,
              'line_spacing': 1.0, 'align': align},
             {'text': unit, 'size': unit_size, 'color': unit_color, 'space_before': 4,
              'align': align}],
            wrap=False)


def bullets(slide, x, y, w, h, items, size=18, col='dgray', spacing=1.55):
    """逐行列表：比用全角空格挤在一行更稳，避免断行难看。"""
    textbox(slide, x, y, w, h,
            [{'text': t, 'size': size, 'color': col, 'line_spacing': spacing} for t in items])


def paras(slide, x, y, w, h, items, size=19, col='dgray', spacing=1.55, gap=10):
    """多段落说明：每段独立，段间距 gap。items 可为字符串或 {'t','c','b','size'}。"""
    lines = []
    for i, it in enumerate(items):
        if isinstance(it, str):
            it = {'t': it}
        lines.append({'text': it['t'], 'size': it.get('size', size),
                      'color': it.get('c', col), 'bold': it.get('b', False),
                      'line_spacing': it.get('spacing', spacing),
                      'space_before': 0 if i == 0 else gap})
    textbox(slide, x, y, w, h, lines)


def panel(slide, x, y, w, h, fill='xlgray', pad=36):
    """带内边距的说明面板，返回内容区左上角 (x, y)。"""
    rect(slide, x, y, w, h, fill=fill)
    return x + pad, y + pad
